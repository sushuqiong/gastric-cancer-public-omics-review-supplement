#!/usr/bin/env python3
"""Build v14 (revised) manuscript DOCX."""
import re
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(r"C:\Users\fengq\Desktop\v14_revision")
MD = ROOT / "manuscript" / "v14_source.md"
OUT = ROOT / "manuscript" / "revised_manuscript_v14.docx"


def parse(text):
    text = text.replace("\r\n", "\n")
    lines = text.splitlines()
    parts = {"title": lines[0].strip(), "front": [], "abstract": [], "body": [],
             "decl": [], "refs": [], "tables": {}, "legend_lines": []}
    mode = "front"
    i = 1
    while i < len(lines):
        line = lines[i].rstrip()
        if line == "## Abstract": mode = "abstract"; i += 1; continue
        if line == "## 1. Introduction": mode = "body"; i += 1; continue
        if line == "## Declarations": mode = "decl"; i += 1; continue
        if line == "## References": mode = "refs"; i += 1; continue
        if line == "## Figure Legends": mode = "legends"; i += 1; continue

        m = re.match(r"\*?\*?Table (\d+)\.\s*(.*)", line)
        if m:
            num = int(m.group(1))
            title = re.sub(r"\*\*", "", m.group(2)).strip().rstrip(".")
            rows, i = [], i + 1
            while i < len(lines) and (lines[i].startswith("|") or not lines[i].strip()):
                if lines[i].startswith("|"):
                    cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                    if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                        rows.append(cells)
                i += 1
            parts["tables"][num] = (title, rows)
            continue

        if mode == "refs":
            if re.match(r"^\d+\.\s+", line): parts["refs"].append(line)
        elif mode == "legends":
            if line.strip(): parts["legend_lines"].append(line)
        elif line.strip():
            parts[mode].append(line)
        i += 1
    return parts


def setup(doc):
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(2.5)
    s.left_margin = s.right_margin = Cm(2.5)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)
    st.paragraph_format.line_spacing = 1.5
    for h in ("Heading 1", "Heading 2"):
        hs = doc.styles[h]
        hs.font.name = "Times New Roman"
        hs.font.size = Pt(12 if h == "Heading 1" else 11)
        hs.font.bold = True
        hs.font.color.rgb = None
    p = s.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for el, attr in (("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "end")):
        e = OxmlElement(el)
        if attr: e.set(qn("w:fldCharType"), attr)
        else: e.set(qn("xml:space"), "preserve"); e.text = " PAGE "
        run._r.append(e)


def add_lines(doc, lines):
    for line in lines:
        t = line.strip()
        if not t: continue
        if t.startswith("## "): doc.add_heading(t[3:], level=1)
        elif t.startswith("### "): doc.add_heading(t[4:], level=2)
        elif t.startswith("**") and "**" in t[2:]:
            doc.add_paragraph(t.replace("**", "").strip())
        else: doc.add_paragraph(t)


def main():
    parts = parse(MD.read_text(encoding="utf-8"))
    doc = Document(); setup(doc)

    # Title page
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(parts["title"]); r.bold = True; r.font.size = Pt(14)
    for line in parts["front"]:
        doc.add_paragraph(line)
    doc.add_page_break()

    # Abstract
    doc.add_heading("Abstract", level=1)
    for line in parts["abstract"]:
        if line.startswith("**Keywords:") or line.startswith("Keywords:"):
            pp = doc.add_paragraph(); pp.add_run("Keywords: ").bold = True
            pp.add_run(re.sub(r"\*+Keywords:?\*+", "", line).strip())
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        else:
            doc.add_paragraph(line)
    doc.add_page_break()

    # Body
    add_lines(doc, parts["body"]); doc.add_page_break()

    # Declarations
    doc.add_heading("Declarations", level=1)
    for line in parts["decl"]:
        if line == "## Declarations": continue
        if line.startswith("### "): doc.add_heading(line[4:], level=2)
        else: doc.add_paragraph(line)
    doc.add_page_break()

    # References
    doc.add_heading("References", level=1)
    for line in parts["refs"]: doc.add_paragraph(line)
    doc.add_page_break()

    # Tables
    for num in sorted(parts["tables"]):
        title, rows = parts["tables"][num]
        pp = doc.add_paragraph(); pp.add_run(f"Table {num}. {title}").bold = True
        if rows:
            t = doc.add_table(rows=1, cols=len(rows[0])); t.style = "Table Grid"
            for j, v in enumerate(rows[0]):
                c = t.rows[0].cells[j]; c.text = v
                for r in c.paragraphs[0].runs: r.bold = True; r.font.size = Pt(9)
            for row in rows[1:]:
                cells = t.add_row().cells
                for j, v in enumerate(row[:len(cells)]): cells[j].text = v
            for r_ in t.rows:
                for c in r_.cells:
                    for para in c.paragraphs:
                        para.paragraph_format.line_spacing = 1.0
                        for run in para.runs: run.font.size = Pt(9)
        if num != max(parts["tables"]): doc.add_page_break()
    doc.add_page_break()

    # Figure legends
    doc.add_heading("Figure Legends", level=1)
    for line in parts["legend_lines"]:
        doc.add_paragraph(line)

    doc.save(OUT)
    txt = "\n".join(p.text for p in Document(OUT).paragraphs)
    assert "**" not in txt, "literal ** leaked"
    assert "|---" not in txt, "stray separator leaked"
    print(f"OK  {OUT.name}")
    print(f"    tables={len(parts['tables'])} refs={len(parts['refs'])} legends={len(parts['legend_lines'])}")


if __name__ == "__main__":
    main()
